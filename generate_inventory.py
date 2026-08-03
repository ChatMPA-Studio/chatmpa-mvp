import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Skill Inventory"

# ── Colors ────────────────────────────────────────────────────────────────────
HEADER_BG    = "1F4E79"
HEADER_FG    = "FFFFFF"
SEC_ERDDAP   = "D6E4F0"
SEC_LTEM     = "D5E8D4"
SEC_CONAP    = "FFF2CC"
SEC_MULTI    = "F3E5F5"
SEC_SHARED   = "EDEDED"

STATUS_TEST    = "FFF2CC"   # implemented, not yet tested in Claude
STATUS_PENDING = "FCE4D6"   # not yet implemented
STATUS_DONE    = "E2EFDA"   # implemented and tested

thin   = Side(style="thin", color="AAAAAA")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

NA = "N/A"

# ── Scale notes (conceptual, not code) ───────────────────────────────────────
ERDDAP_SCALE = (
    "Current: local = the MPA polygon; regional = the Large Marine Ecosystem "
    "(LME) that contains the MPA. Both are fixed geographic boundaries used to "
    "clip satellite data.\n\n"
    "Planned (pending coastal grid): local = coastal grid cells assigned to "
    "that specific MPA; regional = grid cells belonging to the corresponding "
    "fishing management region. Both will be extracted from the same coastal "
    "grid using different spatial filters, replacing the current dual-polygon approach."
)

CONAP_CPUE_SCALE = (
    "Current: local = the nearest commercial fishing office to the MPA "
    "(a nominal administrative proxy, not a spatial boundary); "
    "regional = the entire state where the MPA is located.\n\n"
    "Planned (pending coastal grid): local = georeferenced landing sites "
    "within coastal grid cells assigned to the MPA; regional = landing "
    "sites aggregated by fishing management region. True geographic "
    "assignment will replace administrative proxies."
)

LFO_SCALE = (
    "Operates across all CONAPESCA offices nationally — no local/regional "
    "distinction within the skill itself. Its output (the office-to-region "
    "lookup table) is what defines the regional scale used by other fishing "
    "pressure skills."
)

LTEM_SCALE = (
    "Scale is fixed by the location of LTEM monitoring sites (permanent "
    "transects). The coastal grid does not change the skill's internal logic; "
    "it only provides a shared spatial reference when comparing LTEM outputs "
    "with satellite or fisheries data."
)

PPR_RATIO_SCALE = (
    "Current: combines fishing pressure at the office level (administrative "
    "proxy for the MPA) with ocean productivity over the MPA polygon — "
    "two spatially mismatched units.\n\n"
    "Planned (pending coastal grid): both fishing demand and ecosystem "
    "production will be resolved over the same coastal grid cells assigned "
    "to the MPA, making the ratio spatially consistent."
)

INTERDATABASE_INHERITED = (
    "Scale is inherited from the per-database skills that feed this skill. "
    "No independent spatial logic."
)

PFE_SCALE = (
    "This skill is grid-native by design: the method (Johnson et al. 2017) "
    "distributes fishing effort over equal-area coastal cells from the start. "
    "Local and regional outputs are obtained by filtering those cells by MPA "
    "or fishing management region, respectively."
)

# ── Rows ──────────────────────────────────────────────────────────────────────
# (category, database, skill_name, description, output, repository,
#  scales, status, technical_notes, row_fill, status_fill)

rows = [
    # ── per-database / ERDDAP ─────────────────────────────────────────────────
    (
        "per-database",
        "ERDDAP\n(OISST / NOAA)",
        "erddap-sst-anomaly",
        (
            "Assess thermal conditions and warming trends at a marine protected "
            "area and its surrounding Large Marine Ecosystem, reporting annual SST "
            "anomalies at local and regional scales. Fires on questions about ocean "
            "warming near an MPA, thermal stress on reef communities, El Niño/La "
            "Niña impacts on a specific protected area, or how local sea surface "
            "temperature compares to the regional trend."
        ),
        "Annual mean SST (°C) and mean SST anomaly (°C) per scale, with pixel coverage percentage per year.",
        "per-database/erddap-sst-anomaly",
        ERDDAP_SCALE,
        "Implemented — not yet tested in Claude",
        "Dataset: OISST 0.25°, daily, 1981–present. Baseline: NOAA 1971–2000 climatology. Min. coverage: 50% pixel-days/year.",
        SEC_ERDDAP, STATUS_TEST
    ),
    (
        "per-database",
        "ERDDAP\n(MODIS Aqua)",
        "erddap-chlorophyll",
        (
            "Assess primary productivity and phytoplankton trends at a marine "
            "protected area and its surrounding Large Marine Ecosystem, reporting "
            "annual chlorophyll-a geometric means and anomalies at local and regional "
            "scales. Fires on questions about ocean productivity near an MPA, "
            "upwelling strength, food availability for fish communities, nutrient "
            "conditions, or how local chlorophyll compares to regional patterns."
        ),
        "Annual geometric mean of Chl-a (mg/m³) per scale, with valid-composite coverage percentage per year.",
        "per-database/erddap-chlorophyll",
        ERDDAP_SCALE,
        "Implemented — not yet tested in Claude",
        "Dataset: erdMH1chla8day_R202SQ, 4 km, 8-day composites, 2003–present. Min. coverage: 30%.",
        SEC_ERDDAP, STATUS_TEST
    ),
    (
        "per-database",
        "ERDDAP\n(MODIS Aqua)",
        "erddap-pp",
        (
            "Assess primary productivity levels and trends at a marine protected "
            "area and its surrounding Large Marine Ecosystem, reporting annual net "
            "primary production geometric means and anomalies at local and regional "
            "scales. Fires on questions about ocean productivity near an MPA, "
            "phytoplankton growth efficiency, carbon flux into the food web, or "
            "how local production compares to regional patterns."
        ),
        "Annual geometric mean of NPP (mgC/m²/day) per scale, with valid-composite coverage percentage per year.",
        "per-database/erddap-pp",
        ERDDAP_SCALE,
        "Implemented — not yet tested in Claude",
        "Dataset: erdMH1pp8day, 4 km, 8-day composites, 2003–present. Min. coverage: 30%.",
        SEC_ERDDAP, STATUS_TEST
    ),
    # ── per-database / LTEM ───────────────────────────────────────────────────
    (
        "per-database",
        "LTEM",
        "ltem-nrsi-index",
        (
            "Assess the trophic health of a reef and answer whether it is healthy, "
            "degraded, or recovering, from the LTEM database. Fires on questions "
            "about reef trophic balance, apex-predator recovery, overfishing "
            "signatures, or 'is this reef healthy' — by computing the Normalized "
            "Reef State Index (NRSI) from the relative biomass of upper, lower, "
            "and consumer trophic levels."
        ),
        "NRSI per reef-year with 95% bootstrap confidence intervals. Qualitative trophic state classification.",
        "per-database/ltem-nrsi-index",
        LTEM_SCALE,
        "Implemented — not yet tested in Claude",
        "Standard + conditional NRSI formula. Bootstrap n=100, seed=42. Aggregation: transect → reef-year.",
        SEC_LTEM, STATUS_TEST
    ),
    # ── per-database / CONAPESCA ──────────────────────────────────────────────
    (
        "per-database",
        "CONAPESCA",
        "conapesca-cpue",
        (
            "Estimate fishing pressure on a target species from CONAPESCA landing "
            "records, by computing CPUE (catch per unit effort) as a time series "
            "at regional and local scales. Fires on questions about how fishing "
            "pressure has changed over time for a given species in a given state, "
            "how pressure near a marine protected area compares with the state-level "
            "trend, or which fleet type (industrial vs. artisanal) drives landings."
        ),
        "Annual CPUE time series (kg/effective fishing day) by species, fleet (MAYORES / MENORES) and scale, with effort quality flags.",
        "per-database/conapesca-cpue",
        CONAP_CPUE_SCALE,
        "Implemented — not yet tested in Claude",
        "Analysis unit: trip notice (folio). Excludes folios with generic date flags when computing effort.",
        SEC_CONAP, STATUS_TEST
    ),
    (
        "per-database",
        "CONAPESCA",
        "conapesca-ppr",
        (
            "Estimate the ecological footprint of a fishery by calculating the "
            "Primary Production Required (PPR) to sustain observed catches at a "
            "marine protected area or state. Fires on questions about how much "
            "primary production or phytoplankton the fishery consumes, what the "
            "trophic footprint of a species or fleet is, or whether fishing pressure "
            "is ecologically costly relative to the trophic level of the target species."
        ),
        "PPR (kg C/year) per year, species, fleet and scale (local / regional).",
        "per-database/conapesca-ppr",
        CONAP_CPUE_SCALE,
        "Implemented — not yet tested in Claude",
        "Formula: Pauly & Christensen (1995). Default TE=0.1, WW-to-C=9 (user-overridable). Does not use effort fields.",
        SEC_CONAP, STATUS_TEST
    ),
    (
        "per-database",
        "CONAPESCA",
        "conapesca-lfo-ca",
        (
            "Assess whether commercial fisheries offices in Mexico show distinct "
            "spatial structure based on the composition and volume of their landings. "
            "Fires when the user asks whether landing offices can be grouped into "
            "fishing regions, whether there is a spatial gradient in catch composition "
            "across offices, or as a prerequisite to regionalizing fishing pressure "
            "for downstream analyses. Must be run before conapesca-lfo-regions."
        ),
        "CA scores per office, variance explained per axis, WK and KL quality indices for a range of k values, spatial correlation of CA axis 2 with latitude/longitude.",
        "per-database/conapesca-lfo-ca",
        LFO_SCALE,
        "Implemented — not yet tested in Claude",
        "Ref: Erisman et al. (2011) ICES J. Mar. Sci. 68(3). DCA + CA on office × resource-group matrix.",
        SEC_CONAP, STATUS_TEST
    ),
    (
        "per-database",
        "CONAPESCA",
        "conapesca-lfo-regions",
        (
            "Cluster CONAPESCA fishing offices into discrete fishing regions based "
            "on the similarity of their catch composition, and produce a lookup "
            "table assigning each office to a region. Fires after the user has "
            "reviewed the CA ordination and quality indices from conapesca-lfo-ca "
            "and has decided on the number of regions. The output lookup is consumed "
            "by downstream skills that require regional fishing pressure. "
            "Must be preceded by conapesca-lfo-ca."
        ),
        "Canonical lookup table: office → region ID (numbered north to south). Cluster centroids in CA space and within-cluster variance.",
        "per-database/conapesca-lfo-regions",
        LFO_SCALE,
        "Implemented — not yet tested in Claude",
        "k-means seed=42, nstart=25, iter.max=100. k chosen by user after reviewing WK/KL from lfo-ca.",
        SEC_CONAP, STATUS_TEST
    ),
    (
        "per-database",
        "CONAPESCA",
        "conapesca-pfe",
        (
            "Estimate the spatial distribution of small-scale (artisanal) fishing "
            "effort in the Gulf of California using the method of Johnson et al. "
            "(2017), combining vessel counts from CONAPESCA landing records, "
            "coastal population data from INEGI, and kernel density estimation "
            "over a coastal grid."
        ),
        "Predicted Fishing Effort (PFE) per coastal grid cell: geometric mean of observed and population-predicted vessels. Catch prediction per cell.",
        "per-database/conapesca-pfe",
        PFE_SCALE,
        "Not yet implemented",
        "CONTEXT.md only — no SKILL.md or skill.R. Prerequisites: georeferenced landing ports, validated allometric relationship, coastal grid. Ref: Johnson et al. (2017) PLoS ONE 12(4).",
        SEC_CONAP, STATUS_PENDING
    ),
    # ── interdatabase ─────────────────────────────────────────────────────────
    (
        "interdatabase",
        "LTEM + CONAPESCA\n+ ERDDAP",
        "eda-multivariate",
        (
            "Explore the structure of a multivariate annual time series before "
            "committing to a specific analysis. Computes summary statistics, "
            "normality checks, temporal trends, pairwise correlations, and outlier "
            "detection across all variables simultaneously. Fires on questions about "
            "data structure, whether variables are correlated, whether there are "
            "temporal trends worth noting, or when the user wants to understand "
            "the data before choosing an analysis."
        ),
        "Structured R list with: descriptive statistics per variable (n, mean, median, sd, skewness), Shapiro-Wilk results, temporal trend tests, pairwise Pearson/Spearman correlations, outlier flags. Returns structured objects only — rendering is handled by the LLM/orchestrator, not by this skill.",
        "interdatabase/eda-multivariate",
        INTERDATABASE_INHERITED,
        "Implemented — not yet tested in Claude",
        "Input: aligned table from temporal_align(). No internal transformations. All variables treated symmetrically. Sections are selectable.",
        SEC_MULTI, STATUS_TEST
    ),
    (
        "interdatabase",
        "LTEM + CONAPESCA\n+ ERDDAP",
        "reef-drivers-lm",
        (
            "Identify which environmental and fishing pressure factors are associated "
            "with reef trophic health at a marine protected area, by fitting a linear "
            "model of NRSI against the ecological footprint of fishing (%PPR — fraction "
            "of local primary production consumed by the fishery, per species) and "
            "thermal anomaly (SST). Fires on questions about what drives reef condition, "
            "whether fishing demand or temperature has a stronger association with reef "
            "health, or which factors explain year-to-year variation in reef trophic "
            "state. [Planned version — current implementation uses CPUE and Chl-a; "
            "see GitHub issue #6 for the pending refactor.]"
        ),
        "OLS coefficients with 95% CI and p-values, R² and adjusted R², F-statistic, VIF table, fitted values and residuals per year, Durbin-Watson statistic, ACF of residuals, and number of reefs surveyed per year. All returned as structured R objects (lists/data.frames), not rendered output.",
        "interdatabase/reef-drivers-lm",
        INTERDATABASE_INHERITED,
        "Implemented (current version) — refactor pending (GitHub issue #6)",
        "Current: NRSI ~ CPUE × species × fleet + SST_anom + Chl_anom. Planned: NRSI ~ %PPR × species + SST_anom (Chl removed — already in PPR denominator). Blocked by erddap-pp (#2) and ppr-ratio (#3).",
        SEC_MULTI, STATUS_TEST
    ),
    (
        "interdatabase",
        "CONAPESCA + ERDDAP",
        "ppr-ratio",
        (
            "Estimate what fraction of the ocean's primary production is consumed "
            "by a fishery at a marine protected area, by combining the Primary "
            "Production Required (PPR) by the fishery with the actual net primary "
            "production of the MPA. Fires on questions about the ecological footprint "
            "of fishing relative to ecosystem productivity, whether a fishery is "
            "sustainable given local ocean conditions, or how fishing pressure scales "
            "with primary production across years."
        ),
        "%PPR per year, species and fleet. Number of years excluded due to insufficient satellite coverage.",
        "interdatabase/ppr-ratio",
        PPR_RATIO_SCALE,
        "Implemented — not yet tested in Claude",
        "Sustainability reference: <5% light pressure, >25–35% heavy exploitation (Pauly & Christensen 1995). Requires MPA area in km².",
        SEC_MULTI, STATUS_TEST
    ),
    (
        "interdatabase",
        "Any OLS fitting skill",
        "model-diagnostics",
        (
            "Compute and summarise standard OLS regression diagnostics for any "
            "fitted model that follows the minimal output interface (fitted values, "
            "residuals, formula, n observations). Fires when the user asks whether "
            "model assumptions hold, whether residuals are autocorrelated, or whether "
            "the regression is valid before reporting results. Currently wired to "
            "reef-drivers-lm as the only fitting skill, but the logic is fully generic."
        ),
        "Shapiro-Wilk test on residuals, Durbin-Watson statistic with lag-1 autocorrelation, ACF and PACF of residuals (lags 1–3), White-type heteroscedasticity test. All returned as structured R objects, not rendered output.",
        "interdatabase/model-diagnostics",
        INTERDATABASE_INHERITED,
        "Implemented — not yet tested in Claude",
        "Operates only on residuals and fitted values — no lm object or raw data required. All diagnostics are informative only; the skill does not conclude model adequacy.",
        SEC_MULTI, STATUS_TEST
    ),
    # ── shared ────────────────────────────────────────────────────────────────
    (
        "shared",
        "Any skill with\nan OLS model",
        "results-report",
        (
            "Summarise the numerical results of a fitted regression model and its "
            "diagnostics into a structured, objective statistical report. Fires when "
            "the user asks for the model results to be reported, written up, or "
            "summarised — or when the analysis pipeline is ready to move from "
            "computation to output. Does NOT interpret the ecology; it only "
            "describes what the model says."
        ),
        "Structured R list with: coefficient table with 95% CI, model fit summary (R², AIC, BIC, F, RMSE), data provenance (n years, excluded years, formula), and diagnostic flags if model-diagnostics output is provided. Returns numbers and flags only — rendering (tables, figures, markdown) is handled by Stage 5 of ORCHESTRATION, not by this skill.",
        "shared/results-report",
        NA,
        "Implemented — not yet tested in Claude",
        "Reusable across any fitting skill that follows the declared minimal interface. Produces content (numbers), not format — the LLM/orchestrator decides how to render it.",
        SEC_SHARED, STATUS_TEST
    ),
]

# ── Headers ───────────────────────────────────────────────────────────────────
headers = [
    "Category",
    "Database",
    "Skill name",
    "Description",
    "Output",
    "Repository",
    "Scale\n(current → planned)",
    "Status",
    "Technical notes",
]

header_font  = Font(name="Calibri", bold=True, color=HEADER_FG, size=11)
header_fill  = PatternFill("solid", fgColor=HEADER_BG)
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

for col_idx, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.font      = header_font
    cell.fill      = header_fill
    cell.alignment = header_align
    cell.border    = border

ws.row_dimensions[1].height = 34

# ── Data rows ─────────────────────────────────────────────────────────────────
data_align = Alignment(vertical="top", wrap_text=True)

for row_idx, r in enumerate(rows, start=2):
    values      = r[:9]
    row_fill    = r[9]
    status_fill = r[10]

    for col_idx, val in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.alignment = data_align
        cell.border    = border
        cell.fill = PatternFill("solid", fgColor=(status_fill if col_idx == 8 else row_fill))

    ws.row_dimensions[row_idx].height = 110

# ── Column widths ─────────────────────────────────────────────────────────────
widths = [13, 16, 26, 58, 48, 34, 52, 28, 50]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A2"

# ── Legend sheet ──────────────────────────────────────────────────────────────
ws2 = wb.create_sheet("Legend")
legend = [
    ("Row color",    "Meaning"),
    (SEC_LTEM,       "per-database / LTEM"),
    (SEC_ERDDAP,     "per-database / ERDDAP"),
    (SEC_CONAP,      "per-database / CONAPESCA"),
    (SEC_MULTI,      "interdatabase"),
    (SEC_SHARED,     "shared"),
    ("",             ""),
    ("Status color", "Meaning"),
    (STATUS_TEST,    "Implemented — not yet tested in Claude"),
    (STATUS_PENDING, "Not yet implemented"),
    (STATUS_DONE,    "Implemented and tested in Claude"),
]

ws2.column_dimensions["A"].width = 18
ws2.column_dimensions["B"].width = 46

for i, (color_or_label, label) in enumerate(legend, start=1):
    ca = ws2.cell(row=i, column=1, value="")
    cb = ws2.cell(row=i, column=2, value=label)
    if color_or_label in (SEC_LTEM, SEC_ERDDAP, SEC_CONAP, SEC_MULTI, SEC_SHARED,
                          STATUS_TEST, STATUS_PENDING, STATUS_DONE):
        ca.fill = PatternFill("solid", fgColor=color_or_label)
    elif color_or_label in ("Row color", "Status color"):
        ca.value = color_or_label
        ca.font  = Font(bold=True)
        cb.font  = Font(bold=True)
    ca.border = border
    cb.border = border

# ── Save ──────────────────────────────────────────────────────────────────────
out_path = r"C:\Users\carol\OneDrive\Documentos\CBMC\Projects\chatmpa-mvp\inventario_skills_chatmpa.xlsx"
wb.save(out_path)
print(f"Saved: {out_path}")
